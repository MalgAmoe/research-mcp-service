from __future__ import annotations

import asyncio
import hashlib
import re
import time
from datetime import datetime, timezone
from typing import Any
from urllib.parse import urljoin, urlparse

import httpx
from ddgs import DDGS

from .config import settings
from .errors import GatewayError
from .storage import artifact_store, redis_store
from .models import (
    Artifacts,
    BatchFetchRequest,
    BatchFetchResponse,
    BrowserProfile,
    ConvertDocumentRequest,
    ConvertedDocument,
    Document,
    DocumentMetadata,
    ExtractionAttempt,
    ReadStrategy,
    ReadUrlRequest,
    ResearchFailure,
    ResearchSource,
    ResourceType,
    RssFeed,
    RssFeedEntry,
    RssFetchResponse,
    SearchResult,
    SearchAttempt,
    SearchWebRequest,
    SearchWebResponse,
)


def _domain(url: str) -> str | None:
    parsed = urlparse(url)
    return parsed.hostname or None


def _doc_id(url: str, content: str = "") -> str:
    h = hashlib.sha256((url + "\n" + content[:2048]).encode()).hexdigest()[:16]
    return f"doc_{h}"


def _content_hash(content: str) -> str:
    return "sha256:" + hashlib.sha256(content.encode()).hexdigest()


def _excerpt(text: str, limit: int = 500) -> str:
    compact = re.sub(r"\s+", " ", text).strip()
    return compact[:limit]


def _word_count(text: str) -> int:
    return len(re.findall(r"\w+", text))


def _search_published_at(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None


def _read_metadata_extra(strategy: ReadStrategy, browser_profile: BrowserProfile, cache_hit: bool) -> dict[str, Any]:
    extra: dict[str, Any] = {"cache_hit": cache_hit}
    if strategy == ReadStrategy.browser:
        extra["browser_profile"] = browser_profile.value
    return extra


def _metadata_with_cache_hit(doc: Document, cache_hit: bool) -> DocumentMetadata:
    extra = dict(doc.metadata.extra)
    extra["cache_hit"] = cache_hit
    return doc.metadata.model_copy(update={"extra": extra})


async def search_web(req: SearchWebRequest, *, provider: str | None = None) -> SearchWebResponse:
    provider = provider or settings.search_provider
    if provider not in {'auto', 'ddgs', 'searxng'}:
        raise GatewayError('search_provider_invalid', 'Unknown search provider', status_code=400)
    attempts: list[SearchAttempt] = []

    async def attempt(name, function):
        try:
            response = await function(req)
            attempts.append(SearchAttempt(provider=name, status=response.status,
                result_count=len(response.results), engine_failures=response.engine_failures))
            return response
        except Exception as exc:
            # Preserve routing evidence, never provider exception text or private URLs.
            attempts.append(SearchAttempt(provider=name, status='error',
                error_code=exc.code if isinstance(exc, GatewayError) else 'provider_failed'))
            return None

    primary = None
    result = None
    if provider in {'auto', 'searxng'}:
        primary = await attempt('searxng', _search_searxng)
        if primary is not None and (primary.results or provider == 'searxng'):
            result = primary
    if result is None and provider in {'auto', 'ddgs'}:
        fallback = await attempt('ddgs', _search_ddgs)
        result = fallback if fallback is not None and fallback.results else (primary or fallback)
    if result is None:
        raise GatewayError('search_failed', 'Search providers are unavailable', status_code=502,
            attempts=[item.model_dump() for item in attempts])
    result.attempts = attempts
    return result


async def _search_ddgs(req: SearchWebRequest) -> SearchWebResponse:
    timelimit = {
        "day": "d",
        "week": "w",
        "month": "m",
        "year": "y",
    }.get(req.freshness.value)
    query = _build_search_query(req)

    def run() -> list[dict[str, Any]]:
        # DDGS limits concurrent providers based on max_results. Requesting a
        # wider internal pool makes metasearch genuinely multi-provider; only
        # the caller's requested number is returned below.
        internal_limit = max(req.limit, 30)
        last_error: Exception | None = None
        for _ in range(3):
            try:
                client = DDGS(timeout=15)
                results = client.text(
                    query,
                    region=settings.ddgs_region,
                    safesearch="moderate",
                    timelimit=timelimit,
                    max_results=internal_limit,
                    page=req.page,
                    backend=settings.ddgs_backend,
                )
                if results:
                    return results
            except Exception as exc:
                last_error = exc
        if last_error:
            raise last_error
        return []

    raw_results = await asyncio.to_thread(run)
    results: list[SearchResult] = []
    for item in raw_results:
        url = item.get("href") or item.get("url")
        if not url or re.search(r"https?://(?:www\.)?bing\.com/aclick", url):
            continue
        domain = _domain(url)
        if not _search_domain_allowed(domain, req):
            continue
        resource_type, detected_suffix = _detect_resource_type(url, None)
        results.append(
            SearchResult(
                title=item.get("title") or url,
                url=url,
                snippet=item.get("body"),
                source="ddgs",
                rank=len(results) + 1,
                published_at=_search_published_at(item.get("date")),
                engines=[],  # Configured backends are not verified per-result provenance.
                category=req.category,
                domain=domain,
                resource_type=resource_type,
                detected_suffix=detected_suffix,
            )
        )
        if len(results) >= req.limit:
            break
    return SearchWebResponse(
        query=req.query,
        results=results,
        status="ok" if results else "no_results",
        provider="ddgs",
    )


async def _search_searxng(req: SearchWebRequest) -> SearchWebResponse:
    params: dict[str, Any] = {
        "q": _build_search_query(req),
        "format": "json",
        "pageno": req.page,
        "categories": req.category,
    }
    if req.language:
        params["language"] = req.language
    if req.freshness != "any":
        params["time_range"] = req.freshness.value

    try:
        async with httpx.AsyncClient(timeout=20.0) as client:
            response = await client.get(f"{settings.searxng_url.rstrip('/')}/search", params=params)
            response.raise_for_status()
            payload = response.json()
    except httpx.HTTPError as exc:
        raise GatewayError(
            "search_failed",
            "Could not query SearXNG",
            status_code=502,
            query=req.query,
            cause=str(exc),
        ) from exc
    except ValueError as exc:
        raise GatewayError(
            "search_bad_response",
            "SearXNG returned an invalid response",
            status_code=502,
            query=req.query,
        ) from exc

    raw_results = payload.get("results", [])
    engine_failures = [
        [str(part) for part in failure]
        for failure in payload.get("unresponsive_engines", [])
        if isinstance(failure, (list, tuple))
    ]
    working_engines: set[str] = set()
    results: list[SearchResult] = []
    for item in raw_results:
        url = item.get("url")
        if not url:
            continue
        domain = _domain(url)
        if not _search_domain_allowed(domain, req):
            continue
        engines = [str(engine) for engine in item.get("engines", [])]
        working_engines.update(engines)
        resource_type, detected_suffix = _detect_resource_type(url, item.get("content_type"))
        results.append(
            SearchResult(
                title=item.get("title") or url,
                url=url,
                snippet=item.get("content"),
                source="searxng",
                rank=len(results) + 1,
                published_at=_search_published_at(item.get("publishedDate")),
                engines=engines,
                category=item.get("category"),
                domain=domain,
                resource_type=resource_type,
                detected_suffix=detected_suffix,
            )
        )
        if len(results) >= req.limit:
            break

    if results:
        status = "degraded" if engine_failures or len(working_engines) < 2 else "ok"
    else:
        status = "unavailable" if engine_failures else "no_results"
    return SearchWebResponse(
        query=req.query,
        results=results,
        status=status,
        provider="searxng",
        engine_failures=engine_failures,
    )


def _search_domain_allowed(domain: str | None, req: SearchWebRequest) -> bool:
    host = (domain or '').lower().rstrip('.')
    def matches(value):
        value = value.lower().rstrip('.')
        return host == value or host.endswith('.' + value)
    return (not req.include_domains or any(matches(value) for value in req.include_domains)) and not any(
        matches(value) for value in req.exclude_domains)


def _build_search_query(req: SearchWebRequest) -> str:
    query = req.query
    if req.include_domains:
        clauses = ' OR '.join(f'site:{domain}' for domain in req.include_domains)
        query += ' ' + (f'({clauses})' if len(req.include_domains) > 1 else clauses)
    for domain in req.exclude_domains:
        query += f' -site:{domain}'
    return query


async def read_url(req: ReadUrlRequest) -> Document:
    url = str(req.url)
    _raise_if_read_url_document(url, None)
    if not req.force_refresh:
        cached = await redis_store.get_by_url(url)
        cache_strategy_matches = bool(
            cached
            and (
                cached.strategy_used == req.strategy
                or (
                    req.strategy == ReadStrategy.auto
                    and cached.strategy_used in {
                        ReadStrategy.light,
                        ReadStrategy.standard,
                        ReadStrategy.browser,
                    }
                )
            )
        )
        if cached and cached.strategy_used == ReadStrategy.standard and cached.method != "playwright+markdownify":
            cache_strategy_matches = False
        if cache_strategy_matches and cached and cached.strategy_used == ReadStrategy.browser:
            if cached.metadata.extra.get("browser_profile") != req.browser_profile.value:
                cached = None
                cache_strategy_matches = False
        if cache_strategy_matches and cached:
            if req.include_markdown and cached.artifacts.markdown:
                markdown = artifact_store.get_text_uri(cached.artifacts.markdown)
                return cached.model_copy(update={"markdown": markdown, "metadata": _metadata_with_cache_hit(cached, True)})
            return cached.model_copy(update={"markdown": None, "metadata": _metadata_with_cache_hit(cached, True)})

    title, markdown, method, strategy_used, attempts, raw_html = await _read_with_strategy(req)
    word_count = _word_count(markdown)

    profile_part = req.browser_profile.value if strategy_used == ReadStrategy.browser else ""
    doc_id = _doc_id(f"{url}:{strategy_used}:{profile_part}", markdown)
    now = datetime.now(timezone.utc)

    try:
        raw_uri = None
        if raw_html is not None:
            raw_uri = artifact_store.put_text(
                settings.minio_bucket_raw,
                f"{doc_id}.html",
                raw_html,
                "text/html; charset=utf-8",
            )
        markdown_uri = artifact_store.put_text(
            settings.minio_bucket_processed,
            f"{doc_id}.md",
            markdown,
            "text/markdown; charset=utf-8",
        )
    except Exception as exc:
        raise GatewayError(
            "artifact_store_failed",
            "Could not store page artifacts",
            status_code=502,
            url=url,
            cause=str(exc),
        ) from exc

    stored = Document(
        doc_id=doc_id,
        url=url,
        canonical_url=url,
        title=title,
        markdown=markdown,
        excerpt=_excerpt(markdown),
        retrieved_at=now,
        method=method,
        strategy_requested=req.strategy,
        strategy_used=strategy_used,
        attempts=attempts if req.include_diagnostics else [],
        status="ready",
        artifacts=Artifacts(raw_html=raw_uri, markdown=markdown_uri),
        metadata=DocumentMetadata(
            domain=_domain(url),
            content_hash=_content_hash(markdown),
            word_count=word_count,
            extra=_read_metadata_extra(strategy_used, req.browser_profile, cache_hit=False),
        ),
    )

    try:
        await redis_store.save_document(stored)
    except Exception as exc:
        raise GatewayError(
            "metadata_store_failed",
            "Could not store document metadata",
            status_code=502,
            url=url,
            doc_id=stored.doc_id,
            cause=str(exc),
        ) from exc
    if not req.include_markdown:
        return stored.model_copy(update={"markdown": None})
    return stored


def _detect_resource_type(url: str, content_type: str | None = None) -> tuple[ResourceType, str | None]:
    suffix = _suffix_for_document(url, content_type)
    if suffix in (".pdf", ".docx", ".pptx", ".txt", ".md"):
        return ResourceType.document, suffix.lstrip(".")
    if suffix in (".xlsx", ".csv", ".json", ".parquet", ".zip"):
        return ResourceType.file, suffix.lstrip(".")
    if suffix in (".html", ".htm"):
        return ResourceType.webpage, suffix.lstrip(".")
    return ResourceType.unknown, None


def _raise_if_read_url_document(url: str, content_type: str | None) -> None:
    resource_type, detected_suffix = _detect_resource_type(url, content_type)
    if resource_type in (ResourceType.document, ResourceType.file) and detected_suffix in ("pdf", "docx", "pptx", "xlsx", "csv", "json", "parquet", "zip"):
        suggested_tool = "convert_document" if detected_suffix in ("pdf", "docx", "pptx") else None
        raise GatewayError(
            "file_url_not_supported_by_read_url",
            "Use /convert_document for supported documents; data/archive files are not handled by /read_url",
            status_code=400,
            url=url,
            content_type=content_type,
            detected_suffix=detected_suffix,
            suggested_tool=suggested_tool,
        )


async def _fetch_html(url: str) -> str:
    if settings.fetch_proxy:
        from research_mcp_service.worker_io import download
        response = await download(url, max_bytes=5 * 1024 * 1024)
        _raise_if_read_url_document(url, response.headers.get("content-type"))
        return response.text
    try:
        async with httpx.AsyncClient(follow_redirects=True, timeout=20.0) as client:
            response = await client.get(url, headers={"User-Agent": "local-web-research-gateway/0.1"})
            response.raise_for_status()
            _raise_if_read_url_document(url, response.headers.get("content-type"))
            return response.text
    except httpx.HTTPStatusError as exc:
        raise GatewayError(
            "fetch_bad_status",
            "URL returned an unsuccessful HTTP status",
            status_code=502,
            url=url,
            http_status=exc.response.status_code,
        ) from exc
    except httpx.HTTPError as exc:
        raise GatewayError(
            "fetch_failed",
            "Could not fetch URL",
            status_code=502,
            url=url,
            cause=str(exc),
        ) from exc


def _extract_title(html: str) -> str | None:
    match = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
    if not match:
        return None
    return re.sub(r"\s+", " ", match.group(1)).strip()


async def _read_with_strategy(
    req: ReadUrlRequest,
) -> tuple[str | None, str, str, ReadStrategy, list[ExtractionAttempt], str | None]:
    if settings.worker_socket:
        from research_mcp_service.worker_client import get_worker
        data = await get_worker(settings.worker_socket).run('read', req.model_dump(mode='json'))
        return (data['title'], data['markdown'], data['method'], ReadStrategy(data['strategy_used']),
                [ExtractionAttempt.model_validate(item) for item in data['attempts']], data['raw_html'])
    url = str(req.url)
    attempts: list[ExtractionAttempt] = []

    if req.strategy == ReadStrategy.auto:
        # Start with the cheapest reader and escalate only when extraction is
        # unavailable or too sparse to be useful. Keep the caller-selected
        # browser profile for the final rendered fallback.
        for strategy in (
            ReadStrategy.light,
            ReadStrategy.standard,
            ReadStrategy.browser,
        ):
            candidate = req.model_copy(update={"strategy": strategy})
            try:
                title, markdown, method, strategy_used, candidate_attempts, raw_html = (
                    await _read_with_strategy(candidate)
                )
                attempts.extend(candidate_attempts)
                if _word_count(markdown) >= 30:
                    return title, markdown, method, strategy_used, attempts, raw_html
                if attempts:
                    attempts[-1] = attempts[-1].model_copy(
                        update={
                            "status": "failed",
                            "error": "too little extractable text",
                        }
                    )
            except GatewayError as exc:
                attempts.append(
                    ExtractionAttempt(
                        tool=strategy.value,
                        status="failed",
                        error=f"{exc.code}: {exc.message}",
                    )
                )

        raise GatewayError(
            "auto_extraction_failed",
            "No automatic extraction strategy returned usable text",
            status_code=502,
            url=url,
            attempts=[attempt.model_dump() for attempt in attempts],
        )

    if req.strategy == ReadStrategy.light:
        html = await _fetch_html(url)
        title = _extract_title(html)
        started = time.perf_counter()
        markdown, method = _extract_markdown(html, url)
        duration_ms = int((time.perf_counter() - started) * 1000)
        attempts.append(
            ExtractionAttempt(
                tool=method,
                status="success",
                word_count=_word_count(markdown),
                duration_ms=duration_ms,
            )
        )
        return title, markdown, method, ReadStrategy.light, attempts, html

    if req.strategy == ReadStrategy.standard:
        markdown, title, method, attempt, raw_html = await _extract_with_playwright(
            url, structured=True
        )
        attempts.append(attempt)
        return title, markdown, method, ReadStrategy.standard, attempts, raw_html

    if req.strategy == ReadStrategy.browser:
        markdown, title, method, attempt, rendered_html = await _extract_with_browser(url, req.browser_profile)
        attempts.append(attempt)
        return title, markdown, method, ReadStrategy.browser, attempts, rendered_html

    raise GatewayError(
        "unsupported_strategy",
        "This read strategy is not supported",
        status_code=400,
        url=url,
        strategy=req.strategy.value,
    )


async def _extract_with_browser(
    url: str,
    profile: BrowserProfile,
) -> tuple[str, str | None, str, ExtractionAttempt, str]:
    if profile == BrowserProfile.standard:
        return await _extract_with_playwright(url)
    if profile == BrowserProfile.stealth:
        return await _extract_with_camoufox(url)
    raise GatewayError(
        "unsupported_browser_profile",
        "This browser profile is not implemented",
        status_code=400,
        url=url,
        browser_profile=profile.value,
    )


async def _extract_with_playwright(url: str, *, structured: bool = False) -> tuple[str, str | None, str, ExtractionAttempt, str]:
    try:
        from playwright.async_api import async_playwright  # type: ignore
    except Exception as exc:
        raise GatewayError(
            "browser_dependency_missing",
            "Playwright is not installed",
            status_code=500,
            url=url,
            browser_profile="standard",
            cause=str(exc),
        ) from exc

    started = time.perf_counter()
    browser = None
    try:
        async with async_playwright() as p:
            proxy_options = {"proxy": {"server": settings.fetch_proxy, "bypass": "<-loopback>"}} if settings.fetch_proxy else {}
            browser = await p.chromium.launch(headless=True, **proxy_options)
            page = await browser.new_page()
            response = await page.goto(url, wait_until="domcontentloaded", timeout=30000)
            if response is not None and response.status >= 400:
                raise GatewayError('browser_bad_status', 'URL returned an unsuccessful HTTP status',
                                   status_code=502, http_status=response.status)
            try:
                await page.wait_for_load_state("networkidle", timeout=5000)
            except Exception:
                pass
            title = await page.title()
            html = await page.content()
            final_url = page.url
            await browser.close()
        if structured:
            markdown = _structured_markdown(html, final_url)
            extraction_method = "markdownify"
        else:
            markdown, extraction_method = _extract_markdown(html, url)
        method = f"playwright+{extraction_method}"
        return (
            markdown,
            title or None,
            method,
            ExtractionAttempt(
                tool="playwright",
                status="success",
                word_count=_word_count(markdown),
                duration_ms=int((time.perf_counter() - started) * 1000),
            ),
            html,
        )
    except GatewayError:
        raise
    except Exception as exc:
        if browser is not None:
            try:
                await browser.close()
            except Exception:
                pass
        raise GatewayError(
            "browser_extraction_failed",
            "Could not render URL with Playwright",
            status_code=502,
            url=url,
            browser_profile="standard",
            cause=str(exc),
        ) from exc


async def _extract_with_camoufox(url: str) -> tuple[str, str | None, str, ExtractionAttempt, str]:
    try:
        from camoufox.async_api import AsyncCamoufox  # type: ignore
    except Exception as exc:
        raise GatewayError(
            "browser_dependency_missing",
            "Camoufox is not installed",
            status_code=500,
            url=url,
            browser_profile="stealth",
            cause=str(exc),
        ) from exc

    started = time.perf_counter()
    try:
        proxy_options = {"proxy": {"server": settings.fetch_proxy},
                         "firefox_user_prefs": {"network.proxy.allow_hijacking_localhost": True,
                                                "network.proxy.no_proxies_on": ""}} if settings.fetch_proxy else {}
        async with AsyncCamoufox(headless=True, block_images=True, block_webrtc=True, **proxy_options) as browser:
            page = await browser.new_page()
            response = await page.goto(url, wait_until="domcontentloaded", timeout=30000)
            if response is not None and response.status >= 400:
                raise GatewayError('browser_bad_status', 'URL returned an unsuccessful HTTP status',
                                   status_code=502, http_status=response.status)
            try:
                await page.wait_for_load_state("networkidle", timeout=5000)
            except Exception:
                pass
            title = await page.title()
            html = await page.content()
        markdown, extraction_method = _extract_markdown(html, url)
        method = f"camoufox+{extraction_method}"
        return (
            markdown,
            title or None,
            method,
            ExtractionAttempt(
                tool="camoufox",
                status="success",
                word_count=_word_count(markdown),
                duration_ms=int((time.perf_counter() - started) * 1000),
            ),
            html,
        )
    except GatewayError:
        raise
    except Exception as exc:
        raise GatewayError(
            "browser_extraction_failed",
            "Could not render URL with Camoufox",
            status_code=502,
            url=url,
            browser_profile="stealth",
            cause=str(exc),
        ) from exc


def _structured_markdown(html: str, url: str) -> str:
    """Keep technical structure; original HTML is stored separately by read_url."""
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    soup = BeautifulSoup(html, "html.parser")
    base = soup.find("base", href=True)
    base_url = urljoin(url, base["href"]) if base else url
    for element in soup.select(
        'script, style, noscript, nav, aside, footer, '
        '[role="navigation"], [role="complementary"], [role="contentinfo"]'
    ):
        element.decompose()
    for element in soup.find_all(["a", "img"]):
        attribute = "href" if element.name == "a" else "src"
        if element.get(attribute):
            element[attribute] = urljoin(base_url, element[attribute])
    return markdownify(str(soup.body or soup), heading_style="ATX").strip()


def _extract_markdown(html: str, url: str) -> tuple[str, str]:
    try:
        import trafilatura  # type: ignore

        extracted = trafilatura.extract(html, url=url, output_format="markdown")
        if extracted:
            return extracted, "trafilatura"
    except Exception:
        pass

    # Last-resort text extraction when trafilatura returns no usable content.
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text, "html_text_fallback"


async def batch_fetch(req: BatchFetchRequest) -> BatchFetchResponse:
    sources: list[ResearchSource] = []
    failed_urls: list[ResearchFailure] = []
    lock = asyncio.Lock()
    semaphore = asyncio.Semaphore(req.max_concurrency)

    async def fetch_one(url: Any) -> None:
        url_text = str(url)
        async with semaphore:
            try:
                doc = await read_url(
                    ReadUrlRequest(
                        url=url,
                        strategy=req.strategy,
                        browser_profile=req.browser_profile,
                        force_refresh=req.force_refresh,
                        include_markdown=req.include_markdown,
                        include_diagnostics=req.include_diagnostics,
                    )
                )
                source = ResearchSource(
                    doc_id=doc.doc_id,
                    title=doc.title,
                    url=doc.url,
                    excerpt=doc.excerpt,
                    retrieved_at=doc.retrieved_at,
                    method=doc.method,
                    strategy_used=doc.strategy_used,
                    attempts=doc.attempts,
                    markdown=doc.markdown if req.include_markdown else None,
                )
                async with lock:
                    sources.append(source)
            except GatewayError as exc:
                async with lock:
                    failed_urls.append(
                        ResearchFailure(
                            url=url_text,
                            code=exc.code,
                            message=exc.message,
                            details=exc.context,
                        )
                    )
            except Exception as exc:
                async with lock:
                    failed_urls.append(
                        ResearchFailure(
                            url=url_text,
                            code="read_failed",
                            message=str(exc),
                        )
                    )

    await asyncio.gather(*(fetch_one(url) for url in req.urls))
    order = {str(url): idx for idx, url in enumerate(req.urls)}
    sources.sort(key=lambda source: order.get(source.url, 10**9))
    failed_urls.sort(key=lambda failure: order.get(failure.url, 10**9))
    return BatchFetchResponse(sources=sources, failed_urls=failed_urls)


async def convert_document(req: ConvertDocumentRequest) -> ConvertedDocument:
    from .document_io import prepare_document
    url = str(req.url)
    data, content_type, suffix, markdown, conversion_metadata = await prepare_document(req)
    doc_id = _doc_id(url, data[:2048].hex())

    try:
        original_uri = artifact_store.put_bytes(
            settings.minio_bucket_documents,
            f"{doc_id}{suffix}",
            data,
            content_type or "application/octet-stream",
        )
    except Exception as exc:
        raise GatewayError(
            "artifact_store_failed",
            "Could not store source document",
            status_code=502,
            url=url,
            cause=str(exc),
        ) from exc

    markdown_uri = artifact_store.put_text(
        settings.minio_bucket_processed,
        f"{doc_id}.md",
        markdown,
        "text/markdown; charset=utf-8",
    )

    metadata = DocumentMetadata(
        domain=_domain(url),
        content_hash=_content_hash(markdown),
        word_count=_word_count(markdown),
        extra={"content_type": content_type, "source_suffix": suffix, "strategy": req.strategy.value,
               **conversion_metadata},
    )
    artifacts = Artifacts(raw_html=original_uri, markdown=markdown_uri)
    method = _document_method_for_suffix(suffix)

    # Store converted documents in the same document cache so /documents/{doc_id}
    # and the Pi web_document tool can retrieve them later.
    await redis_store.save_document(
        Document(
            doc_id=doc_id,
            url=url,
            canonical_url=url,
            title=None,
            markdown=markdown,
            excerpt=_excerpt(markdown),
            method=method,
            strategy_requested=None,
            strategy_used=None,
            attempts=[],
            status="ready",
            artifacts=artifacts,
            metadata=metadata,
        )
    )

    converted = ConvertedDocument(
        doc_id=doc_id,
        url=url,
        title=None,
        markdown=markdown if req.include_markdown else None,
        excerpt=_excerpt(markdown),
        method=method,
        artifacts=artifacts,
        metadata=metadata,
    )
    return converted


def _suffix_for_document(url: str, content_type: str | None) -> str:
    parsed_path = urlparse(url).path
    for suffix in (".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".json", ".parquet", ".zip", ".html", ".htm", ".md", ".txt"):
        if parsed_path.lower().endswith(suffix):
            return suffix
    if content_type:
        lowered = content_type.lower()
        if "pdf" in lowered:
            return ".pdf"
        if "wordprocessingml" in lowered or "msword" in lowered:
            return ".docx"
        if "presentationml" in lowered or "powerpoint" in lowered:
            return ".pptx"
        if "spreadsheetml" in lowered or "excel" in lowered:
            return ".xlsx"
        if "csv" in lowered:
            return ".csv"
        if "json" in lowered:
            return ".json"
        if "parquet" in lowered:
            return ".parquet"
        if "zip" in lowered:
            return ".zip"
        if "markdown" in lowered:
            return ".md"
        if lowered.startswith("text/"):
            return ".txt"
    return ".bin"


def _document_method_for_suffix(suffix: str) -> str:
    return {
        ".pdf": "pypdfium2",
        ".docx": "python-docx",
        ".pptx": "python-pptx",
        ".txt": "text",
        ".md": "text",
        ".html": "trafilatura",
        ".htm": "trafilatura",
    }.get(suffix, "unsupported")


async def _convert_document_content(data: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        return _convert_pdf_with_pypdfium(data)
    if suffix in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    if suffix in (".html", ".htm"):
        html = data.decode("utf-8", errors="replace")
        markdown, _method = _extract_markdown(html, "")
        return markdown
    if suffix == ".docx":
        return _convert_docx(data)
    if suffix == ".pptx":
        return _convert_pptx(data)
    raise ValueError(f"unsupported document suffix: {suffix}")


def _convert_pdf_with_pypdfium(data: bytes) -> str:
    import pypdfium2 as pdfium  # type: ignore

    pdf = pdfium.PdfDocument(data)
    chunks: list[str] = []
    for index, page in enumerate(pdf, start=1):
        textpage = page.get_textpage()
        text = textpage.get_text_range().strip()
        if text:
            chunks.append(f"## Page {index}\n\n{text}")
    return "\n\n".join(chunks).strip()


def _convert_docx(data: bytes) -> str:
    from io import BytesIO
    from docx import Document as DocxDocument  # type: ignore

    doc = DocxDocument(BytesIO(data))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts).strip()


def _convert_pptx(data: bytes) -> str:
    from io import BytesIO
    from pptx import Presentation  # type: ignore

    prs = Presentation(BytesIO(data))
    slides: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                lines.append(text.strip())
        if lines:
            slides.append(f"## Slide {index}\n\n" + "\n\n".join(lines))
    return "\n\n".join(slides).strip()



async def get_document(doc_id: str, include_markdown: bool = True) -> Document | None:
    doc = await redis_store.get_document(doc_id)
    if not doc:
        return None
    if include_markdown and doc.artifacts.markdown:
        markdown = artifact_store.get_text_uri(doc.artifacts.markdown)
        return doc.model_copy(update={"markdown": markdown})
    return doc.model_copy(update={"markdown": None})


async def list_recent_documents(limit: int = 20, include_markdown: bool = False) -> list[Document]:
    docs = await redis_store.list_recent_documents(limit=limit)
    if include_markdown:
        hydrated: list[Document] = []
        for doc in docs:
            if doc.artifacts.markdown:
                markdown = artifact_store.get_text_uri(doc.artifacts.markdown)
                hydrated.append(doc.model_copy(update={"markdown": markdown}))
            else:
                hydrated.append(doc)
        return hydrated
    return [doc.model_copy(update={"markdown": None}) for doc in docs]


# ── RSS helpers ────────────────────────────────────────────────────────────

import xml.etree.ElementTree as ET  # noqa: E402


async def _fetch_parsed_feed(feed_url: str) -> dict[str, Any]:
    if settings.worker_socket:
        from research_mcp_service.worker_client import get_worker
        return await get_worker(settings.worker_socket).run('feed', {'feed_url': feed_url})
    try:
        async with httpx.AsyncClient(follow_redirects=True, timeout=20.0) as client:
            response = await client.get(feed_url, headers={
                'User-Agent': 'local-web-research-gateway/0.1',
                'Accept': 'application/rss+xml, application/atom+xml, application/xml, text/xml, */*',
            })
            response.raise_for_status()
    except httpx.HTTPError:
        raise GatewayError('rss_fetch_failed', 'Could not fetch RSS feed', status_code=502) from None
    try:
        return _parse_rss_feed(response.text, feed_url)
    except Exception:
        raise GatewayError('rss_parse_failed', 'Could not parse RSS feed', status_code=502) from None


def _feed_key(feed_url: str) -> str:
    return hashlib.sha256(feed_url.encode()).hexdigest()[:16]


def _entry_id(feed_url: str, entry_link: str | None, entry_title: str | None, published_at: datetime | None) -> str:
    parts = [feed_url, entry_link or "", entry_title or "", str(published_at) if published_at else ""]
    return "rss_" + hashlib.sha256("|".join(parts).encode()).hexdigest()[:16]


async def rss_subscribe(feed_url: str) -> RssFeed:
    key = _feed_key(feed_url)
    existing = await redis_store.get_rss_feed(key)
    if existing:
        return existing
    feed = RssFeed(feed_url=feed_url)
    await redis_store.save_rss_feed(key, feed)
    return feed


async def rss_list_feeds() -> list[RssFeed]:
    return await redis_store.list_rss_feeds()


async def rss_fetch(feed_url: str, max_entries: int = 20, force_refresh: bool = False) -> RssFetchResponse:
    key = _feed_key(feed_url)
    feed = await redis_store.get_rss_feed(key)
    if not feed:
        feed = RssFeed(feed_url=feed_url)

    try:
        parsed = await _fetch_parsed_feed(feed_url)
    except GatewayError as exc:
        feed = feed.model_copy(update={"fetch_status": "error", "fetch_error": exc.message,
                                       "last_fetched_at": datetime.now(timezone.utc)})
        await redis_store.save_rss_feed(key, feed)
        raise

    feed_title = parsed.get("title")
    feed_description = parsed.get("description")
    site_url = parsed.get("link")

    entries: list[RssFeedEntry] = []
    existing_ids = set(await redis_store.get_rss_entry_ids(key))
    new_count = 0

    for item in parsed.get("entries", [])[:max_entries]:
        entry_id = _entry_id(feed_url, item.get("link"), item.get("title"), item.get("published_at"))
        entry = RssFeedEntry(
            entry_id=entry_id,
            feed_url=feed_url,
            title=item.get("title"),
            link=item.get("link"),
            published_at=item.get("published_at"),
            summary=item.get("summary"),
            content=item.get("content"),
        )
        entries.append(entry)
        if entry_id not in existing_ids:
            new_count += 1
            await redis_store.save_rss_entry(key, entry)

    feed = feed.model_copy(update={
        "title": feed.title or feed_title,
        "description": feed.description or feed_description,
        "site_url": feed.site_url or site_url,
        "last_fetched_at": datetime.now(timezone.utc),
        "fetch_status": "ok",
        "fetch_error": None,
        "entry_count": len(await redis_store.get_rss_entry_ids(key)),
    })
    await redis_store.save_rss_feed(key, feed)
    return RssFetchResponse(feed=feed, entries=entries, new_entries=new_count)


async def rss_preview(feed_url: str, max_entries: int = 20) -> RssFetchResponse:
    """Fetch a feed once and return entries without persisting subscription state."""
    parsed = await _fetch_parsed_feed(feed_url)

    feed_title = parsed.get("title")
    feed_description = parsed.get("description")
    site_url = parsed.get("link")

    entries: list[RssFeedEntry] = []
    for item in parsed.get("entries", [])[:max_entries]:
        entry_id = _entry_id(feed_url, item.get("link"), item.get("title"), item.get("published_at"))
        entries.append(
            RssFeedEntry(
                entry_id=entry_id,
                feed_url=feed_url,
                title=item.get("title"),
                link=item.get("link"),
                published_at=item.get("published_at"),
                summary=item.get("summary"),
                content=item.get("content"),
            )
        )

    feed = RssFeed(
        feed_url=feed_url,
        title=feed_title,
        description=feed_description,
        site_url=site_url,
        fetch_status="ok",
        entry_count=len(entries),
    )
    return RssFetchResponse(feed=feed, entries=entries, new_entries=len(entries))


async def rss_get_entries(feed_url: str | None = None, limit: int = 50, include_read: bool = False) -> list[RssFeedEntry]:
    if feed_url:
        key = _feed_key(feed_url)
        return await redis_store.list_rss_entries(key, limit=limit, include_read=include_read)
    all_entries: list[RssFeedEntry] = []
    for feed in await redis_store.list_rss_feeds():
        all_entries.extend(await redis_store.list_rss_entries(_feed_key(feed.feed_url), limit=limit, include_read=include_read))
    all_entries.sort(key=lambda e: e.published_at or e.fetched_at, reverse=True)
    return all_entries[:limit]


async def rss_mark_read(entry_id: str, is_read: bool = True) -> RssFeedEntry | None:
    return await redis_store.update_rss_entry_read(entry_id, is_read)


async def rss_get_entry(entry_id: str) -> RssFeedEntry | None:
    return await redis_store.get_rss_entry(entry_id)


def _parse_rss_feed(xml_text: str, feed_url: str) -> dict[str, Any]:
    root = ET.fromstring(xml_text)

    # Detect Atom vs RSS vs RSS 1.0 (RDF)
    tag = root.tag
    if tag.startswith("{"):
        ns, _, tag_name = tag[1:].partition("}")
    else:
        ns, tag_name = "", tag

    if tag_name == "feed":
        return _parse_atom(root, ns)
    if tag_name == "rss":
        return _parse_rss(root)
    if tag_name == "RDF":
        return _parse_rss10(root, ns)
    # Try channel as root (some malformed feeds)
    channel = root.find("channel")
    if channel is not None:
        return _parse_rss(root)
    raise ValueError(f"Unsupported feed format: root tag '{tag_name}'")


def _parse_rss(root: ET.Element) -> dict[str, Any]:
    channel = root.find("channel")
    if channel is None:
        raise ValueError("RSS feed missing <channel>")

    title = _xml_text(channel, "title")
    link = _xml_text(channel, "link")
    description = _xml_text(channel, "description")

    entries: list[dict[str, Any]] = []
    for item in channel.findall("item"):
        entries.append(_parse_rss_item(item))

    return {"title": title, "link": link, "description": description, "entries": entries}


def _parse_rss10(root: ET.Element, rdf_ns: str) -> dict[str, Any]:
    # RSS 1.0 uses RDF root with channel and item as siblings
    # RSS 1.0 default namespace: http://purl.org/rss/1.0/
    # Dublin Core namespace for dates: http://purl.org/dc/elements/1.1/
    rss10_ns = "http://purl.org/rss/1.0/"
    dc_ns = "http://purl.org/dc/elements/1.1/"

    nsmap: dict[str, str] = {}
    if rdf_ns:
        nsmap["rdf"] = rdf_ns
    nsmap["rss"] = rss10_ns
    nsmap["dc"] = dc_ns

    channel = root.find("rss:channel", nsmap) or root.find("channel", {})
    if channel is None:
        # Some feeds omit channel wrapper — try to find any channel-like element
        channel = root.find(f"{{{rdf_ns}}}channel") if rdf_ns else None
    if channel is None:
        raise ValueError("RSS 1.0 feed missing <channel>")

    def ch_find(tag: str) -> ET.Element | None:
        el = channel.find(f"rss:{tag}", nsmap) if channel is not None else None
        if el is None and channel is not None:
            el = channel.find(tag, {})
        if el is None and channel is not None and rdf_ns:
            el = channel.find(f"{{{rdf_ns}}}{tag}")
        return el

    title = _xml_text_el(ch_find("title"))
    link = _xml_text_el(ch_find("link"))
    description = _xml_text_el(ch_find("description"))

    entries: list[dict[str, Any]] = []
    # Items are direct children of RDF root in RSS 1.0
    items = root.findall("rss:item", nsmap) or root.findall("item", {}) or []
    if rdf_ns and not items:
        items = root.findall(f"{{{rdf_ns}}}item")

    for item in items:
        entries.append(_parse_rss10_item(item, nsmap, rdf_ns))

    return {"title": title, "link": link, "description": description, "entries": entries}


def _parse_rss10_item(item: ET.Element, nsmap: dict[str, str], rdf_ns: str | None) -> dict[str, Any]:
    def find(tag: str) -> ET.Element | None:
        el = item.find(f"rss:{tag}", nsmap)
        if el is None:
            el = item.find(tag, {})
        if el is None and rdf_ns:
            el = item.find(f"{{{rdf_ns}}}{tag}")
        return el

    title = _xml_text_el(find("title"))
    link = _xml_text_el(find("link"))
    description = _xml_text_el(find("description"))

    # RSS 1.0 uses dc:date for publication dates
    dc_date = item.find("dc:date", nsmap)
    if dc_date is None:
        dc_date = item.find("{http://purl.org/dc/elements/1.1/}date")
    published_at = _parse_iso8601_date(_xml_text_el(dc_date)) if dc_date is not None else None

    return {"title": title, "link": link, "summary": description, "content": description, "published_at": published_at}


def _parse_atom(root: ET.Element, ns: str) -> dict[str, Any]:
    nsmap = {"atom": ns} if ns else {}

    def find(path: str) -> ET.Element | None:
        if nsmap:
            return root.find(f"atom:{path}", nsmap)
        return root.find(path)

    title = _xml_text_el(find("title"))
    link_el = find("link")
    link = link_el.get("href") if link_el is not None else None
    subtitle = _xml_text_el(find("subtitle"))

    entries: list[dict[str, Any]] = []
    for entry in root.findall("atom:entry", nsmap) if nsmap else root.findall("entry"):
        entries.append(_parse_atom_entry(entry, nsmap))

    return {"title": title, "link": link, "description": subtitle, "entries": entries}


def _parse_rss_item(item: ET.Element) -> dict[str, Any]:
    title = _xml_text(item, "title")
    link = _xml_text(item, "link")
    description = _xml_text(item, "description")
    pub_date = _xml_text(item, "pubDate")
    published_at = _parse_rfc822_date(pub_date) if pub_date else None
    return {"title": title, "link": link, "summary": description, "content": description, "published_at": published_at}


def _parse_atom_entry(entry: ET.Element, nsmap: dict[str, str]) -> dict[str, Any]:
    def find(path: str) -> ET.Element | None:
        if nsmap:
            return entry.find(f"atom:{path}", nsmap)
        return entry.find(path)

    title = _xml_text_el(find("title"))
    link_el = find("link")
    link = link_el.get("href") if link_el is not None else None
    summary = _xml_text_el(find("summary")) or _xml_text_el(find("content"))
    content = _xml_text_el(find("content")) or summary
    updated = _xml_text_el(find("updated")) or _xml_text_el(find("published"))
    published_at = _parse_iso8601_date(updated) if updated else None
    return {"title": title, "link": link, "summary": summary, "content": content, "published_at": published_at}


def _xml_text(parent: ET.Element, tag: str) -> str | None:
    el = parent.find(tag)
    return _xml_text_el(el)


def _xml_text_el(el: ET.Element | None) -> str | None:
    if el is None:
        return None
    text = "".join(el.itertext()).strip()
    return text if text else None


def _parse_rfc822_date(text: str) -> datetime | None:
    from email.utils import parsedate_to_datetime  # type: ignore
    try:
        return parsedate_to_datetime(text)
    except Exception:
        return None


def _parse_iso8601_date(text: str) -> datetime | None:
    text = text.strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%d"):
        try:
            if fmt.endswith("Z") and text.endswith("Z"):
                text = text[:-1] + "+00:00"
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc if fmt.endswith("%z") or fmt.endswith("Z") else None)
        except ValueError:
            continue
    try:
        from datetime import datetime as _dt  # type: ignore
        return _dt.fromisoformat(text.replace("Z", "+00:00"))
    except Exception:
        return None
